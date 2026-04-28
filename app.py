import os
import json
import re
import base64
import requests
import yaml
from datetime import datetime
from flask import Flask, render_template, request, jsonify, make_response, redirect, url_for
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# 风格配置
STYLES = {
    'warm': {
        'name': '热情阳光',
        'css': 'warm.css',
        'template': 'index.html'
    },
    'emo': {
        'name': 'Emo时刻',
        'css': 'emo.css',
        'template': 'index.html'
    },
    'dark': {
        'name': '暗黑深渊',
        'css': 'dark.css',
        'template': 'index.html'
    }
}

# 默认风格
DEFAULT_STYLE = 'warm'


def get_style_dir(style):
    """获取指定风格的数据目录"""
    return os.path.join(BASE_DIR, 'data', style)


def get_data_dirs(style):
    """获取指定风格的所有数据目录"""
    style_dir = get_style_dir(style)
    return {
        'BASE': style_dir,
        'CHARACTERS': os.path.join(style_dir, 'characters'),
        'GROUPS': os.path.join(style_dir, 'groups'),
        'CHATS_SINGLE': os.path.join(style_dir, 'chats', 'single'),
        'CHATS_GROUP': os.path.join(style_dir, 'chats', 'group'),
        'IMAGES': os.path.join(style_dir, 'images')
    }


def ensure_dirs(style):
    """确保指定风格的所有目录存在"""
    dirs = get_data_dirs(style)
    for d in dirs.values():
        os.makedirs(d, exist_ok=True)


ALLOWED_EXTENSIONS = {'md', 'docx', 'pdf', 'xlsx'}
ALLOWED_IMAGE_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp'}

# 加载配置
CONFIG_PATH = os.path.join(BASE_DIR, 'config.yaml')


def load_config():
    if not os.path.exists(CONFIG_PATH):
        return {'llm_models': [], 'vision_model': None}
    with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
        return yaml.safe_load(f)


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def allowed_image(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_IMAGE_EXTENSIONS


def read_file_content(filepath):
    """读取不同格式文件内容"""
    ext = filepath.rsplit('.', 1)[1].lower()
    content = ''
    try:
        if ext == 'md':
            with open(filepath, 'r', encoding='utf-8') as f:
                content = f.read()
        elif ext == 'docx':
            from docx import Document
            doc = Document(filepath)
            content = '\n'.join([p.text for p in doc.paragraphs])
        elif ext == 'pdf':
            import PyPDF2
            with open(filepath, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                content = '\n'.join([page.extract_text() for page in reader.pages])
        elif ext == 'xlsx':
            import openpyxl
            wb = openpyxl.load_workbook(filepath)
            sheet = wb.active
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append('\t'.join([str(cell) if cell else '' for cell in row]))
            content = '\n'.join(rows)
    except Exception as e:
        content = f'Error reading file: {str(e)}'
    return content


def load_character_data(style, name):
    """加载角色数据，返回(persona, memory, model_name)"""
    dirs = get_data_dirs(style)
    filepath = os.path.join(dirs['CHARACTERS'], f'{name}.md')
    if not os.path.exists(filepath):
        return None, None, None

    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()

    model_name = None
    persona = ''
    memory = ''

    if content.startswith('---\n'):
        parts = content.split('---\n', 2)
        if len(parts) >= 3:
            fm = parts[1]
            body = parts[2]
            for line in fm.split('\n'):
                line = line.strip()
                if line.startswith('model:'):
                    model_name = line.split(':', 1)[1].strip()
            # 从 body 中分离 persona 和 memory
            if '---\n## 历史记忆\n' in body:
                parts2 = body.split('---\n## 历史记忆\n', 1)
                persona = parts2[0].strip()
                memory = parts2[1].strip() if len(parts2) > 1 else ''
            else:
                persona = body.strip()
    else:
        persona = content.strip()

    return persona, memory, model_name


def save_character_data(style, name, persona, memory='', model_name=None):
    """保存角色数据"""
    dirs = get_data_dirs(style)
    filepath = os.path.join(dirs['CHARACTERS'], f'{name}.md')
    content = '---\n'
    if model_name:
        content += f'model: {model_name}\n'
    content += '---\n\n'
    content += persona.strip() + '\n'
    if memory and memory.strip():
        content += '\n---\n## 历史记忆\n\n'
        content += memory.strip() + '\n'
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(content)


def call_llm(messages, model_config):
    """调用LLM模型"""
    url = f"{model_config['base_url'].rstrip('/')}/chat/completions"
    headers = {
        'Content-Type': 'application/json',
        'Authorization': f"Bearer {model_config['api_key']}"
    }
    payload = {
        'model': model_config['model_name'],
        'messages': messages,
        'temperature': 0.7
    }
    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=120)
        resp.raise_for_status()
        result = resp.json()
        return result['choices'][0]['message']['content']
    except Exception as e:
        return f"[错误] 模型调用失败: {str(e)}"


def analyze_image(style, image_path):
    """使用视觉模型分析图片"""
    config = load_config()
    vision_config = config.get('vision_model')
    if not vision_config:
        return "未配置视觉模型，无法识别图片"

    # 读取图片并base64编码
    with open(image_path, 'rb') as f:
        img_data = base64.b64encode(f.read()).decode('utf-8')

    ext = image_path.rsplit('.', 1)[1].lower()
    mime_type = f'image/{ext}' if ext != 'jpg' else 'image/jpeg'

    url = f"{vision_config['base_url'].rstrip('/')}/chat/completions"
    headers = {
        'Content-Type': 'application/json',
        'Authorization': f"Bearer {vision_config['api_key']}"
    }
    payload = {
        'model': vision_config['model_name'],
        'messages': [
            {
                'role': 'user',
                'content': [
                    {'type': 'text', 'text': '请详细描述这张图片的内容，包括物体、场景、文字等所有可见元素。'},
                    {'type': 'image_url', 'image_url': {'url': f'data:{mime_type};base64,{img_data}'}}
                ]
            }
        ],
        'max_tokens': 1000
    }
    try:
        resp = requests.post(url, headers=headers, json=payload, timeout=60)
        resp.raise_for_status()
        result = resp.json()
        return result['choices'][0]['message']['content']
    except Exception as e:
        return f"[图片识别失败: {str(e)}]"


def parse_chat_metadata(filepath):
    """解析聊天文件元数据"""
    metadata = {'type': 'single', 'name': '', 'characters': []}
    if not os.path.exists(filepath):
        return metadata

    with open(filepath, 'r', encoding='utf-8') as f:
        line = f.readline()
        if line.strip() == '---':
            while True:
                line = f.readline()
                if not line or line.strip() == '---':
                    break
                line = line.strip()
                if ':' in line:
                    key, value = line.split(':', 1)
                    key = key.strip()
                    value = value.strip()
                    if key == 'type':
                        metadata['type'] = value
                    elif key == 'name':
                        metadata['name'] = value
                    elif key == 'characters' and value:
                        metadata['characters'] = value.split(',')
    return metadata


def parse_chat_messages(filepath):
    """解析聊天记录为messages格式"""
    messages = []
    if not os.path.exists(filepath):
        return messages

    with open(filepath, 'r', encoding='utf-8') as f:
        # 跳过元数据头
        line = f.readline()
        if line.strip() == '---':
            while True:
                line = f.readline()
                if not line or line.strip() == '---':
                    break

        current_sender = None
        current_content = []
        # 继续处理剩余内容
        while True:
            if not line:
                line = f.readline()
            if not line:
                break
            if line.startswith('**'):
                if current_sender and current_content:
                    role = 'user' if current_sender == '你' else 'assistant'
                    messages.append({
                        'role': role,
                        'content': '\n'.join(current_content).strip()
                    })
                sender_match = re.match(r'\*\*(.+?)\*\*', line)
                if sender_match:
                    current_sender = sender_match.group(1)
                current_content = []
            else:
                if current_sender:
                    current_content.append(line.rstrip('\n'))
            line = None
        if current_sender and current_content:
            role = 'user' if current_sender == '你' else 'assistant'
            messages.append({
                'role': role,
                'content': '\n'.join(current_content).strip()
            })
    return messages


# === 路由 ===

@app.route('/')
def index():
    style = request.args.get('style')
    if not style:
        return render_template('home.html')
    if style not in STYLES:
        style = DEFAULT_STYLE
    ensure_dirs(style)
    # 设置风格cookie
    resp = make_response(render_template('index.html', style=style, style_config=STYLES[style]))
    resp.set_cookie('ai_talk_style', style, max_age=30*24*3600)
    return resp


@app.route('/home')
def home():
    return render_template('home.html')


# 从cookie获取当前风格
def get_current_style():
    style = request.cookies.get('ai_talk_style')
    if not style or style not in STYLES:
        style = DEFAULT_STYLE
    return style


@app.route('/data/<style>/images/<filename>')
def serve_style_image(style, filename):
    if style not in STYLES:
        style = DEFAULT_STYLE
    dirs = get_data_dirs(style)
    from flask import send_from_directory
    return send_from_directory(dirs['IMAGES'], filename)


@app.route('/api/models', methods=['GET'])
def get_models():
    config = load_config()
    return jsonify({
        'llm_models': config.get('llm_models', []),
        'has_vision': config.get('vision_model') is not None
    })


@app.route('/api/characters', methods=['GET'])
def get_characters():
    style = get_current_style()
    dirs = get_data_dirs(style)
    characters = []
    if os.path.exists(dirs['CHARACTERS']):
        for filename in os.listdir(dirs['CHARACTERS']):
            if filename.endswith('.md'):
                name = filename[:-3]
                persona, memory, model_name = load_character_data(style, name)
                characters.append({
                    'name': name,
                    'persona': persona or '',
                    'memory': memory or '',
                    'model': model_name
                })
    return jsonify(characters)


@app.route('/api/characters/<name>', methods=['GET'])
def get_character(name):
    style = get_current_style()
    persona, memory, model_name = load_character_data(style, name)
    if persona is None:
        return jsonify({'error': '角色不存在'}), 404
    return jsonify({
        'name': name,
        'persona': persona or '',
        'memory': memory or '',
        'model': model_name
    })


@app.route('/api/characters', methods=['POST'])
def create_character():
    style = get_current_style()
    ensure_dirs(style)
    dirs = get_data_dirs(style)
    data = request.form
    name = data.get('name', '').strip()
    persona = data.get('persona', '').strip()
    model_name = data.get('model', '').strip() or None

    if not name:
        return jsonify({'error': '角色名称不能为空'}), 400

    # 处理上传的文件作为历史记忆
    file_content = ''
    if 'file' in request.files:
        file = request.files['file']
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            temp_path = os.path.join(dirs['BASE'], filename)
            file.save(temp_path)
            file_content = read_file_content(temp_path)
            os.remove(temp_path)

    save_character_data(style, name, persona, file_content, model_name)
    return jsonify({
        'name': name,
        'persona': persona,
        'memory': file_content,
        'model': model_name
    })


@app.route('/api/characters/<name>', methods=['PUT'])
def update_character(name):
    style = get_current_style()
    dirs = get_data_dirs(style)
    data = request.form
    persona = data.get('persona', '').strip()
    memory = data.get('memory', '').strip()
    model_name = data.get('model', '').strip() or None

    # 处理上传的文件（如果有新文件上传，替换历史记忆）
    file_content = None
    if 'file' in request.files:
        file = request.files['file']
        if file and file.filename and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            temp_path = os.path.join(dirs['BASE'], filename)
            file.save(temp_path)
            file_content = read_file_content(temp_path)
            os.remove(temp_path)

    # 如果没有上传新文件，使用传入的memory
    if file_content is None:
        final_memory = memory
    else:
        final_memory = file_content

    save_character_data(style, name, persona, final_memory, model_name)
    return jsonify({
        'name': name,
        'persona': persona,
        'memory': final_memory,
        'model': model_name
    })


@app.route('/api/characters/<name>', methods=['DELETE'])
def delete_character(name):
    style = get_current_style()
    dirs = get_data_dirs(style)
    filepath = os.path.join(dirs['CHARACTERS'], f'{name}.md')
    if os.path.exists(filepath):
        os.remove(filepath)
        return jsonify({'success': True})
    return jsonify({'error': '角色不存在'}), 404


# === 群聊管理 ===
@app.route('/api/groups', methods=['GET'])
def get_groups():
    style = get_current_style()
    dirs = get_data_dirs(style)
    groups = []
    if os.path.exists(dirs['GROUPS']):
        for filename in sorted(os.listdir(dirs['GROUPS'])):
            if filename.endswith('.json'):
                filepath = os.path.join(dirs['GROUPS'], filename)
                with open(filepath, 'r', encoding='utf-8') as f:
                    group = json.load(f)
                    groups.append(group)
    return jsonify(groups)


@app.route('/api/groups', methods=['POST'])
def create_group():
    style = get_current_style()
    ensure_dirs(style)
    dirs = get_data_dirs(style)
    data = request.json
    name = data.get('name', '').strip()
    characters = data.get('characters', [])

    if not name:
        return jsonify({'error': '群聊名称不能为空'}), 400
    if len(characters) < 2:
        return jsonify({'error': '请至少选择2个角色'}), 400

    group = {
        'id': name,
        'name': name,
        'characters': characters
    }

    os.makedirs(dirs['GROUPS'], exist_ok=True)
    filepath = os.path.join(dirs['GROUPS'], f'{name}.json')
    with open(filepath, 'w', encoding='utf-8') as f:
        json.dump(group, f, ensure_ascii=False, indent=2)

    return jsonify(group)


@app.route('/api/groups/<name>', methods=['DELETE'])
def delete_group(name):
    style = get_current_style()
    dirs = get_data_dirs(style)
    filepath = os.path.join(dirs['GROUPS'], f'{name}.json')
    if os.path.exists(filepath):
        os.remove(filepath)
        return jsonify({'success': True})
    return jsonify({'error': '群聊不存在'}), 404


# === 记忆总结 ===
def collect_character_chats(style, character_name):
    """收集角色相关的所有聊天记录，按天数分类"""
    dirs = get_data_dirs(style)
    chats_by_date = {}

    # 收集单聊
    if os.path.exists(dirs['CHATS_SINGLE']):
        for filename in os.listdir(dirs['CHATS_SINGLE']):
            if filename.startswith(f'{character_name}_') and filename.endswith('.md'):
                # 从文件名提取日期
                date_str = None
                match = re.search(r'_(\d{8})_', filename)
                if match:
                    date_str = match.group(1)
                else:
                    date_str = 'unknown'

                filepath = os.path.join(dirs['CHATS_SINGLE'], filename)
                messages = parse_chat_messages_for_summary(filepath)
                if messages:
                    if date_str not in chats_by_date:
                        chats_by_date[date_str] = []
                    chats_by_date[date_str].append({
                        'type': '单聊',
                        'name': character_name,
                        'date': date_str,
                        'messages': messages
                    })

    # 收集群聊
    if os.path.exists(dirs['CHATS_GROUP']):
        for filename in os.listdir(dirs['CHATS_GROUP']):
            if filename.endswith('.md'):
                # 从文件名提取日期
                date_str = None
                match = re.search(r'_(\d{8})_', filename)
                if match:
                    date_str = match.group(1)
                else:
                    date_str = 'unknown'

                filepath = os.path.join(dirs['CHATS_GROUP'], filename)
                metadata = parse_chat_metadata(filepath)
                if character_name in metadata.get('characters', []):
                    messages = parse_chat_messages_for_summary(filepath)
                    if messages:
                        if date_str not in chats_by_date:
                            chats_by_date[date_str] = []
                        chats_by_date[date_str].append({
                            'type': '群聊',
                            'name': metadata.get('name', filename),
                            'date': date_str,
                            'messages': messages
                        })

    return chats_by_date


def parse_chat_messages_for_summary(filepath):
    """解析聊天记录用于总结"""
    messages = []
    if not os.path.exists(filepath):
        return messages

    with open(filepath, 'r', encoding='utf-8') as f:
        # 跳过元数据头
        line = f.readline()
        if line.strip() == '---':
            while True:
                line = f.readline()
                if not line or line.strip() == '---':
                    break

        current_sender = None
        current_content = []
        while True:
            if not line:
                line = f.readline()
            if not line:
                break
            if line.startswith('**'):
                if current_sender and current_content:
                    messages.append({
                        'sender': current_sender,
                        'content': '\n'.join(current_content).strip()
                    })
                sender_match = re.match(r'\*\*(.+?)\*\*', line)
                if sender_match:
                    current_sender = sender_match.group(1)
                current_content = []
            else:
                if current_sender:
                    current_content.append(line.rstrip('\n'))
            line = None
        if current_sender and current_content:
            messages.append({
                'sender': current_sender,
                'content': '\n'.join(current_content).strip()
            })
    return messages


@app.route('/api/characters/<name>/summarize-memory', methods=['POST'])
def summarize_memory(name):
    style = get_current_style()
    config = load_config()

    # 加载角色
    persona, memory, model_name = load_character_data(style, name)
    if persona is None:
        return jsonify({'error': '角色不存在'}), 404

    # 收集相关聊天（按天分类）
    chats_by_date = collect_character_chats(style, name)
    if not chats_by_date:
        return jsonify({'error': '没有找到相关聊天记录'}), 400

    # 构建总结prompt - 按天组织
    chat_texts = []
    sorted_dates = sorted(chats_by_date.keys())
    for date_str in sorted_dates:
        date_chats = chats_by_date[date_str]
        if date_str != 'unknown':
            try:
                dt = datetime.strptime(date_str, '%Y%m%d')
                display_date = dt.strftime('%Y年%m月%d日')
            except:
                display_date = date_str
        else:
            display_date = '其他'

        chat_text = f'【{display_date}】\n'
        for chat in date_chats:
            chat_text += f'  [{chat["type"]} - {chat["name"]}]\n'
            for msg in chat['messages']:
                chat_text += f'    {msg["sender"]}: {msg["content"]}\n'
        chat_texts.append(chat_text)

    full_prompt = f'''请总结以下角色「{name}」的聊天记录，按天提取关键的记忆点、重要信息、人物关系等。
重点提取与{name}相关的内容。

聊天记录：
{chr(10).join(chat_texts)}

请用简洁的语言总结，保留重要细节。'''

    # 选择模型
    if not model_name and config.get('llm_models'):
        model_name = config['llm_models'][0]['name']

    model_config = None
    for m in config.get('llm_models', []):
        if m['name'] == model_name:
            model_config = m
            break

    if not model_config:
        return jsonify({'error': '未配置模型'}), 400

    # 调用LLM总结
    summary = call_llm([{'role': 'user', 'content': full_prompt}], model_config)

    # 处理原有记忆 - 去重同一天的旧总结
    today_date = datetime.now().strftime('%Y-%m-%d')
    new_memory_parts = []

    if memory and memory.strip():
        # 按分隔符拆分原有记忆
        memory_sections = re.split(r'\n---\n', memory)
        for section in memory_sections:
            section = section.strip()
            if not section:
                continue
            # 检查是否是今天的旧总结
            if f'记忆总结 ({today_date})' in section:
                continue  # 跳过今天的旧总结
            new_memory_parts.append(section)

    # 添加新总结
    new_summary_section = f'## 记忆总结 ({today_date})\n\n{summary}'
    new_memory_parts.append(new_summary_section)

    # 合并所有记忆部分
    new_memory = '\n\n---\n\n'.join(new_memory_parts)

    # 保存回角色文件
    save_character_data(style, name, persona, new_memory, model_name)

    return jsonify({
        'success': True,
        'summary': summary,
        'memory': new_memory
    })


@app.route('/api/chats', methods=['GET'])
def get_chats():
    style = get_current_style()
    dirs = get_data_dirs(style)
    chat_type = request.args.get('type', 'single')
    chats_dir = dirs['CHATS_SINGLE'] if chat_type == 'single' else dirs['CHATS_GROUP']
    chats = []

    if os.path.exists(chats_dir):
        for filename in sorted(os.listdir(chats_dir), reverse=True):
            if filename.endswith('.md'):
                filepath = os.path.join(chats_dir, filename)
                name = filename[:-3]
                match = re.search(r'_(\d{8}_\d{6})$', name)
                if match:
                    timestamp = match.group(1)
                    display_name = name[:-16]
                    try:
                        dt = datetime.strptime(timestamp, '%Y%m%d_%H%M%S')
                        time_str = dt.strftime('%Y-%m-%d %H:%M')
                    except:
                        time_str = timestamp
                else:
                    display_name = name
                    time_str = ''

                # 解析元数据获取角色列表
                metadata = parse_chat_metadata(filepath)
                characters = metadata.get('characters', [])

                # 获取预览
                preview = ''
                with open(filepath, 'r', encoding='utf-8') as f:
                    line = f.readline()
                    if line.strip() == '---':
                        while True:
                            line = f.readline()
                            if not line or line.strip() == '---':
                                break
                    while True:
                        if not line:
                            line = f.readline()
                        if not line:
                            break
                        if line.startswith('**'):
                            preview = line.strip()
                            break
                        line = None

                chats.append({
                    'filename': filename,
                    'name': display_name,
                    'time': time_str,
                    'preview': preview,
                    'type': chat_type,
                    'characters': characters
                })
    return jsonify(chats)


@app.route('/api/chats/<chat_type>/<filename>', methods=['GET'])
def get_chat(chat_type, filename):
    style = get_current_style()
    dirs = get_data_dirs(style)
    chats_dir = dirs['CHATS_SINGLE'] if chat_type == 'single' else dirs['CHATS_GROUP']
    filepath = os.path.join(chats_dir, filename)
    if not os.path.exists(filepath):
        return jsonify({'error': '会话不存在'}), 404

    # 解析元数据
    metadata = parse_chat_metadata(filepath)

    messages = []
    with open(filepath, 'r', encoding='utf-8') as f:
        # 跳过元数据头
        line = f.readline()
        if line.strip() == '---':
            while True:
                line = f.readline()
                if not line or line.strip() == '---':
                    break

        current_sender = None
        current_content = []
        while True:
            if not line:
                line = f.readline()
            if not line:
                break
            if line.startswith('**'):
                if current_sender:
                    messages.append({
                        'sender': current_sender,
                        'content': '\n'.join(current_content).strip()
                    })
                sender_match = re.match(r'\*\*(.+?)\*\* \((\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\)', line)
                if sender_match:
                    current_sender = sender_match.group(1)
                else:
                    current_sender = line.strip('* \n')
                current_content = []
            else:
                if current_sender:
                    current_content.append(line.rstrip('\n'))
            line = None
        if current_sender:
            messages.append({
                'sender': current_sender,
                'content': '\n'.join(current_content).strip()
            })
    return jsonify({'messages': messages, 'metadata': metadata})


@app.route('/api/upload-image', methods=['POST'])
def upload_image():
    style = get_current_style()
    ensure_dirs(style)
    dirs = get_data_dirs(style)
    if 'image' not in request.files:
        return jsonify({'error': '没有上传文件'}), 400

    file = request.files['image']
    if not file or not file.filename:
        return jsonify({'error': '无效的文件'}), 400

    if not allowed_image(file.filename):
        return jsonify({'error': '不支持的图片格式'}), 400

    os.makedirs(dirs['IMAGES'], exist_ok=True)
    ext = file.filename.rsplit('.', 1)[1].lower()
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f'img_{timestamp}.{ext}'
    filepath = os.path.join(dirs['IMAGES'], filename)
    file.save(filepath)

    # 分析图片
    description = analyze_image(style, filepath)

    return jsonify({
        'filename': filename,
        'description': description
    })


@app.route('/api/chats/send', methods=['POST'])
def send_message():
    style = get_current_style()
    ensure_dirs(style)
    dirs = get_data_dirs(style)
    data = request.json
    chat_type = data.get('type', 'single')
    target_name = data.get('target')
    message = data.get('message', '').strip()
    image_desc = data.get('image_description', '')
    image_filename = data.get('image_filename', '')
    character_names = data.get('characters', [])
    is_new = data.get('is_new', False)

    if not message and not image_desc:
        return jsonify({'error': '消息不能为空'}), 400

    config = load_config()

    chats_dir = dirs['CHATS_SINGLE'] if chat_type == 'single' else dirs['CHATS_GROUP']

    if is_new:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'{target_name}_{timestamp}.md'
    else:
        filename = data.get('filename')
        if not filename:
            return jsonify({'error': '缺少filename'}), 400

    filepath = os.path.join(chats_dir, filename)

    # 如果是继续历史群聊且没有角色列表，从元数据读取
    if not is_new and chat_type == 'group' and not character_names:
        metadata = parse_chat_metadata(filepath)
        character_names = metadata.get('characters', [])

    # 构建完整用户消息（包含图片链接和识别内容）
    full_user_message = message
    if image_filename and image_desc:
        image_part = f'![图片](/data/{style}/images/{image_filename})\n[图片内容: {image_desc}]'
        full_user_message = f'{message}\n\n{image_part}' if message else image_part
    elif image_desc:
        full_user_message = f'{message}\n\n[图片内容: {image_desc}]' if message else f'[图片内容: {image_desc}]'

    # 保存用户消息
    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    user_line = f'**你** ({now})\n{full_user_message}\n\n'

    if os.path.exists(filepath):
        with open(filepath, 'a', encoding='utf-8') as f:
            f.write(user_line)
    else:
        # 新文件：写入元数据头
        os.makedirs(chats_dir, exist_ok=True)
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write('---\n')
            f.write(f'type: {chat_type}\n')
            f.write(f'name: {target_name}\n')
            if chat_type == 'group' and character_names:
                f.write(f'characters: {",".join(character_names)}\n')
            f.write('---\n\n')
            f.write(user_line)

    # 生成回复
    replies = []

    if chat_type == 'single':
        # 单聊：调用角色绑定的模型
        # 如果是继续历史聊天，从元数据获取角色名
        if not is_new:
            metadata = parse_chat_metadata(filepath)
            if metadata.get('name'):
                target_name = metadata['name']
        persona, memory, model_name = load_character_data(style, target_name)
        if not model_name and config.get('llm_models'):
            model_name = config['llm_models'][0]['name']

        model_config = None
        for m in config.get('llm_models', []):
            if m['name'] == model_name:
                model_config = m
                break

        if model_config:
            history_messages = parse_chat_messages(filepath)
            system_prompt = f'你是{target_name}。{persona or ""}'
            if memory and memory.strip():
                system_prompt += f'\n\n## 历史记忆\n{memory}'
            system_prompt += f'\n\n请始终以{target_name}的身份和语气回复用户。'
            messages = [{'role': 'system', 'content': system_prompt}] + history_messages
            reply_content = call_llm(messages, model_config)
        else:
            reply_content = f'收到你的消息了："{full_user_message}" (未配置模型)'

        reply_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        reply_line = f'**{target_name}** ({reply_time})\n{reply_content}\n\n'
        with open(filepath, 'a', encoding='utf-8') as f:
            f.write(reply_line)
        replies.append({'sender': target_name, 'content': reply_content, 'time': reply_time})

    else:
        # 群聊：按顺序让每个角色回复
        for char_name in character_names:
            persona, memory, model_name = load_character_data(style, char_name)
            if not model_name and config.get('llm_models'):
                model_name = config['llm_models'][0]['name']

            model_config = None
            for m in config.get('llm_models', []):
                if m['name'] == model_name:
                    model_config = m
                    break

            if model_config:
                history_messages = parse_chat_messages(filepath)
                others = [n for n in character_names if n != char_name]
                system_prompt = f'你是{char_name}。{persona or ""}'
                if memory and memory.strip():
                    system_prompt += f'\n\n## 历史记忆\n{memory}'
                system_prompt += f'\n\n这是一个群聊，其他角色有：{", ".join(others)}。\n请以{char_name}的身份和语气回复。'
                messages = [{'role': 'system', 'content': system_prompt}] + history_messages
                reply_content = call_llm(messages, model_config)
            else:
                reply_content = f'[{char_name}] 收到了消息。'

            reply_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            reply_line = f'**{char_name}** ({reply_time})\n{reply_content}\n\n'
            with open(filepath, 'a', encoding='utf-8') as f:
                f.write(reply_line)
            replies.append({'sender': char_name, 'content': reply_content, 'time': reply_time})

    return jsonify({
        'filename': filename,
        'replies': replies
    })


@app.route('/api/chats/<chat_type>/<filename>', methods=['DELETE'])
def delete_chat(chat_type, filename):
    style = get_current_style()
    dirs = get_data_dirs(style)
    chats_dir = dirs['CHATS_SINGLE'] if chat_type == 'single' else dirs['CHATS_GROUP']
    filepath = os.path.join(chats_dir, filename)
    if os.path.exists(filepath):
        os.remove(filepath)
        return jsonify({'success': True})
    return jsonify({'error': '会话不存在'}), 404


# === 获取当前风格信息 ===
@app.route('/api/style', methods=['GET'])
def get_style_info():
    style = get_current_style()
    return jsonify({
        'style': style,
        'name': STYLES[style]['name'],
        'css': STYLES[style]['css']
    })


if __name__ == '__main__':
    # 确保所有风格的目录都存在
    for s in STYLES:
        ensure_dirs(s)
    app.run(debug=True, port=5002)
